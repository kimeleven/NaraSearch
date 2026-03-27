import datetime
import os
import smtplib
import sys
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

import pandas as pd
import requests
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import config


class BidSearchClient:
    def __init__(self, api_key: str):
        self.api_key = api_key
        self.base_url = config.API_BASE_URL
        self.session = requests.Session()
        self.session.headers.update({"User-Agent": "Mozilla/5.0"})

    def _build_params(self, date_str: str, page_no: int) -> dict:
        return {
            "serviceKey": self.api_key,
            "pageNo": page_no,
            "numOfRows": config.ROWS_PER_PAGE,
            "type": "json",
            "inqryDiv": 1,
            "inqryBgnDt": date_str + "0000",
            "inqryEndDt": date_str + "2359",
        }

    def fetch_page(self, date_str: str, page_no: int) -> list:
        params = self._build_params(date_str, page_no)
        resp = self.session.get(
            self.base_url,
            params=params,
            timeout=config.REQUEST_TIMEOUT_SECONDS,
        )
        resp.raise_for_status()

        content_type = resp.headers.get("Content-Type", "")
        if "json" not in content_type:
            raise ValueError(
                f"JSON 응답이 아닙니다. API 키를 확인하세요.\n"
                f"Content-Type: {content_type}\n"
                f"응답 내용(앞 300자): {resp.text[:300]}"
            )

        data = resp.json()
        header = data.get("response", {}).get("header", {})
        result_code = header.get("resultCode", "")
        if result_code != "00":
            result_msg = header.get("resultMsg", "알 수 없는 오류")
            raise RuntimeError(f"API 오류 [{result_code}]: {result_msg}")

        body = data.get("response", {}).get("body", {})
        items = body.get("items")

        # 결과가 없으면 items가 None, 빈 문자열, 빈 리스트로 옴
        if not items:
            return []

        # 이 API는 items가 바로 list로 옴
        if isinstance(items, list):
            return items
        # 혹시 dict 구조인 경우 (item 키 하위)
        if isinstance(items, dict):
            item = items.get("item", [])
            if isinstance(item, dict):
                return [item]
            return item if item else []
        return []

    def fetch_all(self, date_str: str) -> list:
        all_items = []
        page = 1
        while True:
            print(f"  페이지 {page} 조회 중...", end="\r")
            items = self.fetch_page(date_str, page)
            if not items:
                break
            all_items.extend(items)
            if len(items) < config.ROWS_PER_PAGE:
                break
            page += 1
        print()
        return all_items


class BizInfoClient:
    def __init__(self, api_key: str):
        self.api_key = api_key
        self.session = requests.Session()
        self.session.headers.update({"User-Agent": "Mozilla/5.0"})

    def fetch_page(self, page_no: int) -> list:
        params = {
            "crtfcKey": self.api_key,
            "dataType": "json",
            "pageUnit": config.BIZINFO_ROWS_PER_PAGE,
            "pageIndex": page_no,
        }
        resp = self.session.get(
            config.BIZINFO_API_URL,
            params=params,
            timeout=config.REQUEST_TIMEOUT_SECONDS,
        )
        resp.raise_for_status()
        data = resp.json()
        items = data.get("jsonArray", [])
        if not items:
            print(f"  [기업마당 경고] jsonArray 없음 또는 빈 응답. 응답 키: {list(data.keys())}, 내용 앞부분: {str(data)[:200]}")
        return items

    def fetch_recent(self, days: int = 1) -> list:
        """최근 days일 이내 등록된 공고 수집 (creatPnttm 기준)."""
        cutoff = (datetime.date.today() - datetime.timedelta(days=days)).strftime("%Y-%m-%d")
        result = []
        page = 1
        while True:
            print(f"  [기업마당] 페이지 {page} 조회 중...", end="\r")
            items = self.fetch_page(page)
            if not items:
                break
            found_older = False
            for item in items:
                created = item.get("creatPnttm", "")
                if created and created[:10] >= cutoff:
                    result.append(item)
                elif created and created[:10] < cutoff:
                    found_older = True
                    break
            if found_older or len(items) < config.BIZINFO_ROWS_PER_PAGE:
                break
            page += 1
        print()
        return result


def filter_records(records: list) -> list:
    seen = set()
    result = []
    for r in records:
        name = r.get("bidNtceNm", "")
        task_type = r.get("srvceDivNm", "")
        bid_no = r.get("bidNtceNo", id(r))

        match_a = (task_type in config.ALLOWED_TASK_TYPES
                   and any(kw in name for kw in config.KEYWORDS))
        match_b = any(kw.lower() in name.lower() for kw in config.TABLET_KEYWORDS)

        exclude = any(kw in name for kw in ["대행", "감리", "연구용역"])
        if (match_a or match_b) and bid_no not in seen and not exclude:
            seen.add(bid_no)
            result.append(r)
    return result



def build_bizinfo_dataframe(records: list) -> pd.DataFrame:
    rows = [{field: r.get(field, "") for field in config.BIZINFO_EXCEL_COLUMNS} for r in records]
    df = pd.DataFrame(rows, columns=list(config.BIZINFO_EXCEL_COLUMNS.keys()))
    df.rename(columns=config.BIZINFO_EXCEL_COLUMNS, inplace=True)
    df.sort_values("등록일시", ascending=True, inplace=True)
    df.reset_index(drop=True, inplace=True)
    return df


def build_dataframe(records: list) -> pd.DataFrame:
    rows = []
    for r in records:
        row = {field: r.get(field, "") for field in config.EXCEL_COLUMNS}
        rows.append(row)

    df = pd.DataFrame(rows, columns=list(config.EXCEL_COLUMNS.keys()))
    df.rename(columns=config.EXCEL_COLUMNS, inplace=True)
    df.sort_values("공고일시", ascending=True, inplace=True)
    df.reset_index(drop=True, inplace=True)
    return df


def _style_sheet(ws) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="2F5496")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            (sum(2 if ord(c) > 127 else 1 for c in str(cell.value)) for cell in col_cells if cell.value),
            default=10,
        )
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)


def export_to_excel(df: pd.DataFrame, date_str: str, bizinfo_df: pd.DataFrame = None) -> str:
    display_date = f"{date_str[:4]}-{date_str[4:6]}-{date_str[6:]}"
    filename = f"입찰공고_{display_date}.xlsx"

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="나라장터")
        _style_sheet(writer.sheets["나라장터"])

        if bizinfo_df is not None and not bizinfo_df.empty:
            bizinfo_df.to_excel(writer, index=False, sheet_name="정부지원사업")
            _style_sheet(writer.sheets["정부지원사업"])

    return filename


def send_email(excel_path: str, date_str: str, nara_count: int, bizinfo_count: int = 0) -> None:
    sender = os.environ.get("EMAIL_FROM", "")
    recipients_raw = os.environ.get("EMAIL_TO", "")
    password = os.environ.get("EMAIL_PASSWORD", "")

    if not all([sender, recipients_raw, password]):
        print("이메일 설정(EMAIL_FROM/EMAIL_TO/EMAIL_PASSWORD)이 없어 발송 생략.")
        return

    # 쉼표로 구분된 여러 수신자 지원
    recipients = [r.strip() for r in recipients_raw.split(",") if r.strip()]

    display_date = f"{date_str[:4]}-{date_str[4:6]}-{date_str[6:]}"
    subject = f"[공고검색] {display_date} 나라장터 {nara_count}건 / 기업마당 {bizinfo_count}건"

    body = (
        f"{display_date} 공고 검색 결과입니다.\n\n"
        f"■ 나라장터 입찰공고: {nara_count}건\n"
        f"  - 업무구분: {', '.join(config.ALLOWED_TASK_TYPES)}\n"
        f"  - 키워드: {', '.join(config.KEYWORDS + config.TABLET_KEYWORDS)}\n\n"
        f"■ 기업마당 정부지원사업: {bizinfo_count}건 (최근 1일 신규 공고 전체)\n\n"
        "첨부 엑셀 파일을 확인해 주세요."
    )

    msg = MIMEMultipart()
    msg["From"] = sender
    msg["To"] = ", ".join(recipients)
    msg["Subject"] = subject
    msg.attach(MIMEText(body, "plain", "utf-8"))

    with open(excel_path, "rb") as f:
        attachment = MIMEApplication(f.read(), Name=os.path.basename(excel_path))
        attachment["Content-Disposition"] = f'attachment; filename="{os.path.basename(excel_path)}"'
        msg.attach(attachment)

    with smtplib.SMTP("smtp.gmail.com", 587) as smtp:
        smtp.ehlo()
        smtp.starttls()
        smtp.login(sender, password)
        smtp.sendmail(sender, recipients, msg.as_string())

    print(f"이메일 발송 완료 → {', '.join(recipients)}")


def main():
    yesterday = datetime.date.today() - datetime.timedelta(days=1)
    date_str = yesterday.strftime("%Y%m%d")
    display_date = yesterday.strftime("%Y-%m-%d")

    try:
        api_key = config.get_api_key()
    except RuntimeError as e:
        print(f"[오류] {e}")
        sys.exit(1)

    print(f"날짜 {display_date} 의 입찰공고를 조회합니다...")
    client = BidSearchClient(api_key)

    try:
        raw_records = client.fetch_all(date_str)
    except (requests.HTTPError, requests.Timeout, requests.ConnectionError) as e:
        print(f"[오류] API 호출 실패: {e}")
        sys.exit(1)
    except (ValueError, RuntimeError) as e:
        print(f"[오류] {e}")
        sys.exit(1)

    print(f"총 {len(raw_records)}건 조회됨")

    filtered = filter_records(raw_records)
    print(
        f"필터링 결과: {len(filtered)}건 "
        f"(A: 업무구분({', '.join(config.ALLOWED_TASK_TYPES)}) + "
        f"키워드({', '.join(config.KEYWORDS)}) | "
        f"B: 태블릿({', '.join(config.TABLET_KEYWORDS)}))"
    )

    # ── 기업마당 수집 ──────────────────────────────────────────────
    print(f"\n[기업마당] 최근 1일 정부지원사업 공고를 조회합니다...")
    bizinfo_client = BizInfoClient(config.get_bizinfo_api_key())
    try:
        raw_bizinfo = bizinfo_client.fetch_recent(days=1)
        print(f"[기업마당] {len(raw_bizinfo)}건 조회됨 (최근 1일)")
    except Exception as e:
        print(f"[기업마당 오류] {e}")
        raw_bizinfo = []

    bizinfo_df = build_bizinfo_dataframe(raw_bizinfo) if raw_bizinfo else pd.DataFrame()

    # ── 저장 및 발송 ───────────────────────────────────────────────
    if not filtered and bizinfo_df.empty:
        print("조건에 맞는 공고가 없습니다.")
        return

    df = build_dataframe(filtered) if filtered else pd.DataFrame()
    output_path = export_to_excel(df, date_str, bizinfo_df if not bizinfo_df.empty else None)
    print(f"저장 완료: {output_path}  (나라장터 {len(df)}건 / 기업마당 {len(bizinfo_df)}건)")

    send_email(output_path, date_str, len(df), len(bizinfo_df))


if __name__ == "__main__":
    main()
