"""Daily 2nd-pass filters + combine Excel for 나라장터 / 기업마당 / extra sources.

Filters
-------
나라장터
  - Exclude 수의시담 / 전자시담 / 소액수의 / 수의계약 / 일반경쟁→수의 style awards
  - Sort: 최적가·최저가·규격가격 → 적격심사 → 협상 → 기타

기업마당 (정부지원사업)
  - No extra region filter — keep API results as-is

K-Startup
  - Prefer open (모집중); do not force region filter

과기정통부 (MSIT)
  - Keep yesterday pressDt (already fetched with 1-day lookback edge);
    no 수의 filter
"""

from __future__ import annotations

import argparse
import datetime
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import extra_sources

ROOT = Path("/workspace/nara-search")

# ── 나라장터 ─────────────────────────────────────────────────────────────────

SUUI_EXCLUDE_PATTERNS = [
    re.compile(p)
    for p in [
        r"수의시담",
        r"전자시담",
        r"소액수의",
        r"수의계약",
        r"일반경쟁\s*→\s*수의",
        r"일반경쟁\s*->\s*수의",
        r"수의견적",
    ]
]


def _award_sort_key(method: str) -> tuple[int, str]:
    m = (method or "").strip()
    if any(k in m for k in ("최적가", "최저가", "규격가격")):
        return (0, m)
    if "적격심사" in m:
        return (1, m)
    if "협상" in m:
        return (2, m)
    return (3, m)


def is_suui_excluded(row: pd.Series) -> bool:
    blob = " ".join(
        str(row.get(c, "") or "")
        for c in ("낙찰방법", "낙찰방법세부기준", "입찰방식", "공고명")
    )
    return any(p.search(blob) for p in SUUI_EXCLUDE_PATTERNS)


def filter_nara(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty:
        return df.copy() if df is not None else pd.DataFrame()
    out = df.loc[~df.apply(is_suui_excluded, axis=1)].copy()
    if "낙찰방법" in out.columns:
        out["_sort"] = out["낙찰방법"].map(_award_sort_key)
        out.sort_values("_sort", inplace=True)
        out.drop(columns=["_sort"], inplace=True)
    out.reset_index(drop=True, inplace=True)
    return out


# ── 기업마당 지역 ────────────────────────────────────────────────────────────

ALLOWED_LOCAL_ORGS = {"서울특별시", "경기도"}

# 주관기관이 지방자치단체로 보이면 제외 (서울/경기 제외)
LOCAL_ORG_RE = re.compile(
    r"(광역시|특별자치시|특별자치도|특별시|도)$|통합특별시"
)

# 제목 앞쪽 지역 태그 — 허용/비허용
TITLE_TAG_RE = re.compile(r"^\s*\[([^\]]+)\]")
ALLOWED_TITLE_TAGS = {
    "경기",
    "서울",
    "전국",
    "중앙",
}
# anything else in [] near the start is treated as regional and dropped
# (e.g. [경북], [부산], [충남], [전남광주], [부산ㆍ울산ㆍ경남])


def _org_allowed(org: str) -> bool:
    org = (org or "").strip()
    if not org:
        return True
    if org in ALLOWED_LOCAL_ORGS:
        return True
    # 경기도/서울특별시 already allowed; other *-도 / 광역시 / 특별시 drop
    if org.endswith("도") and org not in ALLOWED_LOCAL_ORGS:
        # 중소벤처기업부 etc. don't end with 도; 충청남도 does
        if "부" in org and not any(x in org for x in ("광역시", "특별시", "특별자치", "통합")):
            # e.g. unlikely; ministries usually end with 부/청/원
            pass
        if re.search(r"(북도|남도|특별자치도)$", org) or org in {
            "강원도",
            "제주도",
        }:
            return False
        if org.endswith("도") and org not in ALLOWED_LOCAL_ORGS:
            # 경기도 allowed above; other 도
            return False
    if LOCAL_ORG_RE.search(org) and org not in ALLOWED_LOCAL_ORGS:
        return False
    if "통합특별시" in org:
        return False
    return True


def _title_region_allowed(title: str) -> bool:
    title = title or ""
    m = TITLE_TAG_RE.match(title)
    if not m:
        return True
    tag = m.group(1).strip()
    # allow if tag is/contains only 경기/서울/전국
    if tag in ALLOWED_TITLE_TAGS:
        return True
    if tag.startswith("경기") or tag.startswith("서울") or tag.startswith("전국"):
        return True
    # multi-region tags like 부산ㆍ울산ㆍ경남 → exclude
    return False


def filter_bizinfo_region(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty:
        return df.copy() if df is not None else pd.DataFrame()
    org_col = "주관기관" if "주관기관" in df.columns else None
    title_col = "공고명" if "공고명" in df.columns else None

    def keep(row: pd.Series) -> bool:
        org = str(row.get(org_col, "") if org_col else "")
        title = str(row.get(title_col, "") if title_col else "")
        if not _org_allowed(org):
            return False
        if not _title_region_allowed(title):
            return False
        return True

    out = df.loc[df.apply(keep, axis=1)].copy()
    out.reset_index(drop=True, inplace=True)
    return out


# ── K-Startup region ─────────────────────────────────────────────────────────

KSTARTUP_ALLOWED_REGIONS = {"전국", "경기", "서울", "중앙"}


def filter_kstartup(df: pd.DataFrame) -> pd.DataFrame:
    if df is None or df.empty:
        return df.copy() if df is not None else pd.DataFrame()
    out = df.copy()
    if "모집중여부" in out.columns:
        out = out.loc[out["모집중여부"].astype(str).str.upper() != "N"]
    # region filter disabled
    out.reset_index(drop=True, inplace=True)
    return out


def filter_msit(df: pd.DataFrame, target_date: datetime.date, lookback_days: int = 1) -> pd.DataFrame:
    if df is None or df.empty:
        return df.copy() if df is not None else pd.DataFrame()
    start = (target_date - datetime.timedelta(days=lookback_days)).isoformat()
    end = target_date.isoformat()
    out = df.copy()
    if "게시일" in out.columns:
        dates = out["게시일"].astype(str).str[:10]
        out = out.loc[(dates >= start) & (dates <= end)]
    out.reset_index(drop=True, inplace=True)
    return out


# ── Excel styling / write ────────────────────────────────────────────────────


def _style_sheet(ws) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="2F5496")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = max(
            (
                sum(2 if ord(c) > 127 else 1 for c in str(cell.value))
                for cell in col_cells
                if cell.value
            ),
            default=10,
        )
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)


def write_combined_excel(
    sheets: dict[str, pd.DataFrame],
    out_path: Path,
) -> Path:
    out_path = Path(out_path)
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            use = df if df is not None else pd.DataFrame()
            use.to_excel(writer, index=False, sheet_name=name)
            _style_sheet(writer.sheets[name])
    return out_path


def load_existing_daily(path: Path) -> dict[str, pd.DataFrame]:
    if not path.exists():
        return {}
    xl = pd.ExcelFile(path)
    return {name: pd.read_excel(path, sheet_name=name) for name in xl.sheet_names}


def build_v2(
    target_date: datetime.date,
    *,
    source_xlsx: Path | None = None,
    out_path: Path | None = None,
) -> tuple[Path, dict[str, int]]:
    date_str = target_date.isoformat()
    if source_xlsx is None:
        source_xlsx = ROOT / f"입찰공고_{date_str}.xlsx"
    if out_path is None:
        out_path = ROOT / f"입찰공고_{date_str}_v2.xlsx"

    existing = load_existing_daily(source_xlsx)
    nara = existing.get("나라장터", pd.DataFrame())
    biz = existing.get("정부지원사업", pd.DataFrame())

    nara_f = filter_nara(nara)
    # 기업마당: 지역 필터 하지 않음 (사용자 요청)
    biz_f = biz.copy() if biz is not None else __import__('pandas').DataFrame()
    if not biz_f.empty:
        biz_f = biz_f.reset_index(drop=True)

    print(f"[filter] 나라장터 {len(nara)} → {len(nara_f)}", flush=True)
    print(f"[filter] 정부지원사업 {len(biz)} → {len(biz_f)}", flush=True)

    extras = extra_sources.fetch_all_extra(target_date)
    ks = filter_kstartup(extras.get("K-Startup", pd.DataFrame()))
    ms = filter_msit(extras.get("과기정통부", pd.DataFrame()), target_date, lookback_days=1)
    # Prefer exact-day MSIT rows; if empty keep lookback set
    if "게시일" in ms.columns and not ms.empty:
        exact = ms.loc[ms["게시일"].astype(str).str[:10] == date_str]
        if not exact.empty:
            ms = exact.reset_index(drop=True)

    print(f"[filter] K-Startup → {len(ks)}", flush=True)
    print(f"[filter] 과기정통부 → {len(ms)}", flush=True)

    sheets = {
        "나라장터": nara_f,
        "정부지원사업": biz_f,
        "K-Startup": ks,
        "과기정통부": ms,
    }
    path = write_combined_excel(sheets, out_path)
    counts = {k: len(v) for k, v in sheets.items()}
    return path, counts


def _sample_titles(df: pd.DataFrame, n: int = 3) -> list[str]:
    if df is None or df.empty:
        return []
    for col in ("공고명", "제목"):
        if col in df.columns:
            return [str(x) for x in df[col].head(n).tolist()]
    return [str(df.iloc[i, 0]) for i in range(min(n, len(df)))]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--date",
        default="2026-09-08",
        help="Target date YYYY-MM-DD (KST yesterday for the daily run)",
    )
    parser.add_argument("--source", default=None, help="Existing daily xlsx path")
    parser.add_argument("--out", default=None, help="Output v2 xlsx path")
    args = parser.parse_args()
    target = datetime.date.fromisoformat(args.date)
    source = Path(args.source) if args.source else None
    out = Path(args.out) if args.out else None
    path, counts = build_v2(target, source_xlsx=source, out_path=out)
    print(f"\nSaved: {path}")
    for name, cnt in counts.items():
        print(f"  {name}: {cnt}")
    # reload for samples
    for name in counts:
        df = pd.read_excel(path, sheet_name=name)
        samples = _sample_titles(df)
        print(f"  samples[{name}]:")
        for s in samples:
            print(f"    - {s}")


if __name__ == "__main__":
    main()
